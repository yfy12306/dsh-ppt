# -*- coding: utf-8 -*-
"""落盘校验 —— PyMuPDF 渲染 PDF，检查文本溢出画布 / 逐页出图供目检。

沿用上游 harness-anything/pptengine/verify.py 的思路（这是它最有价值的一块：
生成完不算完，必须回读 PDF 验证文字有没有跑出画布），并补三项：

1. **重叠检测**：同一页内文本块两两相交面积超阈值即报警（上游只查溢出）。
2. **空页检测**：整页没有任何文本/图像 → 提示可能是渲染失败。
3. **无 PyMuPDF 时优雅降级**：返回 `available=False` 而不是抛异常，
   生成主流程不受影响（上游也是这个策略，保留）。
"""
from __future__ import annotations

import os
from pathlib import Path
from typing import Optional


def _has_fitz() -> bool:
    try:
        import fitz  # noqa: F401
        return True
    except ImportError:
        return False


def _iter_spans(page):
    for b in page.get_text("dict")["blocks"]:
        if b.get("type") != 0:
            continue
        for l in b.get("lines", []):
            for sp in l.get("spans", []):
                t = (sp.get("text") or "").strip()
                if t:
                    yield t, sp["bbox"], sp.get("size", 0)


def _inter_area(a, b) -> float:
    x0 = max(a[0], b[0])
    y0 = max(a[1], b[1])
    x1 = min(a[2], b[2])
    y1 = min(a[3], b[3])
    if x1 <= x0 or y1 <= y0:
        return 0.0
    return (x1 - x0) * (y1 - y0)


def _area(b) -> float:
    return max(0.0, (b[2] - b[0])) * max(0.0, (b[3] - b[1]))


def verify_pdf(pdf_path, width: float = 960.0, height: float = 540.0,
               tol: float = 1.0, overlap_ratio: float = 0.45,
               verbose: bool = True) -> dict:
    """校验 PDF。

    Returns:
        {"available": bool, "pages": int, "overflow": int, "overlaps": int,
         "emptyPages": [..], "issues": [{page,kind,bbox,text}, ..], "pdf": path}
    """
    result = {"available": False, "pages": 0, "overflow": 0, "overlaps": 0,
              "emptyPages": [], "issues": [], "pdf": str(pdf_path)}
    if not _has_fitz():
        result["error"] = "未安装 PyMuPDF（pip install pymupdf），跳过校验"
        if verbose:
            print(f"[verify] {result['error']}")
        return result
    if not os.path.exists(str(pdf_path)):
        result["error"] = f"PDF 不存在：{pdf_path}"
        return result

    import fitz

    result["available"] = True
    doc = fitz.open(str(pdf_path))
    result["pages"] = doc.page_count

    for pi, pg in enumerate(doc, 1):
        W = pg.rect.width or width
        H = pg.rect.height or height
        spans = list(_iter_spans(pg))
        if not spans and not pg.get_images():
            result["emptyPages"].append(pi)
            result["issues"].append({"page": pi, "kind": "empty", "bbox": None, "text": ""})
            if verbose:
                print(f"  p{pi} 空页（无文本无图像）")

        for t, bb, _sz in spans:
            if bb[0] < -tol or bb[1] < -tol or bb[2] > W + tol or bb[3] > H + tol:
                result["overflow"] += 1
                result["issues"].append({"page": pi, "kind": "overflow",
                                         "bbox": [round(v, 1) for v in bb], "text": t[:40]})
                if verbose:
                    print(f"  p{pi} 溢出 [{bb[0]:.0f},{bb[1]:.0f},{bb[2]:.0f},{bb[3]:.0f}] {t[:24]}")

        # 重叠检测：O(n^2)，单页文本块通常 < 60 个，可接受
        if overlap_ratio and overlap_ratio > 0:
            n = len(spans)
            for i in range(n):
                for j in range(i + 1, n):
                    a, b = spans[i][1], spans[j][1]
                    inter = _inter_area(a, b)
                    if inter <= 0:
                        continue
                    smaller = min(_area(a), _area(b))
                    if smaller > 0 and inter / smaller >= overlap_ratio:
                        result["overlaps"] += 1
                        result["issues"].append({
                            "page": pi, "kind": "overlap",
                            "bbox": [round(v, 1) for v in a],
                            "text": f"{spans[i][0][:18]} ⨯ {spans[j][0][:18]}"})
                        if verbose:
                            print(f"  p{pi} 重叠 {spans[i][0][:16]} ⨯ {spans[j][0][:16]}")
    doc.close()

    if verbose:
        print(f"[verify] 页数={result['pages']} 溢出={result['overflow']} "
              f"重叠={result['overlaps']} 空页={len(result['emptyPages'])}")
    return result


def render_pages(pdf_path, out_dir, dpi: int = 110, max_pages: int = 0) -> list[str]:
    """把 PDF 逐页渲染成 PNG，便于人工目检 / 客户端预览。"""
    if not _has_fitz():
        return []
    import fitz

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(str(pdf_path))
    files = []
    for i, pg in enumerate(doc, 1):
        if max_pages and i > max_pages:
            break
        f = out / f"p{i:02d}.png"
        pg.get_pixmap(dpi=dpi).save(str(f))
        files.append(str(f))
    doc.close()
    return files


def verify_pptx(pptx_path) -> dict:
    """不依赖 PDF 的轻量校验：用 python-pptx 读回结构（若可用）。

    用于「保存了但没导 PDF」时至少确认页数与形状数对得上。
    """
    out = {"available": False, "path": str(pptx_path), "slides": 0, "shapes": 0}
    if not os.path.exists(str(pptx_path)):
        out["error"] = f"文件不存在：{pptx_path}"
        return out
    out["size"] = os.path.getsize(str(pptx_path))
    try:
        from pptx import Presentation
    except ImportError:
        out["error"] = "未安装 python-pptx，仅做存在性/大小校验"
        return out
    try:
        prs = Presentation(str(pptx_path))
        out["available"] = True
        out["slides"] = len(prs.slides)
        out["shapes"] = sum(len(s.shapes) for s in prs.slides)
        out["width"] = prs.slide_width
        out["height"] = prs.slide_height
    except Exception as e:  # noqa: BLE001
        out["error"] = f"读取失败：{type(e).__name__}: {e}"
    return out
